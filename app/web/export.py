"""Matrix export (Phase 13): CSV + XLSX (with codebook-colored tag cells), and a
bootstrap parser that learns a tag vocabulary (+ colors) from an old sheet/CSV."""

import csv
import io
from typing import Any, Dict, List, Optional

_BASE_COLS = ["source", "year", "schema"]

_PALETTE = [
    "FFD966", "9FC5E8", "B6D7A8", "EA9999", "D5A6BD",
    "F9CB9C", "B4A7D6", "A2C4C9", "FFE599", "D9D9D9",
]


def _rows_to_table(view: Dict[str, Any]) -> (List[str], List[List[str]]):
    cols = view.get("columns", [])
    header = _BASE_COLS + cols
    table = []
    for r in view.get("rows", []):
        fields = r.get("fields", {})
        row = [
            str(r.get("source", "")),
            str(r.get("year", "") if r.get("year") is not None else ""),
            str(r.get("schema", "")),
        ] + [str(fields.get(c, "")) for c in cols]
        table.append(row)
    return header, table


def to_csv(view: Dict[str, Any]) -> bytes:
    header, table = _rows_to_table(view)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(table)
    return buf.getvalue().encode("utf-8-sig")  # BOM -> clean import to Sheets/Excel


def to_xlsx(view: Dict[str, Any], codebook: Optional[List[Dict[str, Any]]] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header, table = _rows_to_table(view)
    cols = view.get("columns", [])
    tag_col_idx = (_BASE_COLS + cols).index("tag") if "tag" in cols else None

    # tag -> hex color from codebook.
    color_map = {}
    for c in (codebook or []):
        if c.get("tag") and c.get("color"):
            color_map[c["tag"].lower()] = c["color"].lstrip("#")

    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"

    bold = Font(bold=True)
    ws.append(header)
    for cell in ws[1]:
        cell.font = bold

    for row in table:
        ws.append(row)

    # Color the tag cells per codebook.
    if tag_col_idx is not None and color_map:
        for r_i, row in enumerate(table, start=2):
            tag = row[tag_col_idx]
            hexc = color_map.get((tag or "").lower())
            if hexc:
                ws.cell(row=r_i, column=tag_col_idx + 1).fill = PatternFill(
                    start_color=hexc, end_color=hexc, fill_type="solid"
                )

    wrap = Alignment(wrap_text=True, vertical="top")
    for col_cells in ws.columns:
        ws.column_dimensions[col_cells[0].column_letter].width = 28
        for cell in col_cells:
            cell.alignment = wrap

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def bootstrap_codebook_from_bytes(data: bytes, filename: str) -> List[Dict[str, Any]]:
    """Learn tag vocabulary (+ colors from XLSX fills) from an old matrix file.

    Looks for a 'tag' column (case-insensitive). CSV -> tags only (palette colors
    assigned). XLSX -> reads the cell fill color of each tag as its color.
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _bootstrap_xlsx(data)
    return _bootstrap_csv(data)


def _assign_palette(entries: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    for i, e in enumerate(entries):
        if not e.get("color"):
            e["color"] = _PALETTE[i % len(_PALETTE)]
    return entries


def _bootstrap_csv(data: bytes) -> List[Dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    header = [h.strip().lower() for h in rows[0]]
    if "tag" not in header:
        return []
    ti = header.index("tag")
    seen, entries = set(), []
    for r in rows[1:]:
        if ti < len(r):
            tag = (r[ti] or "").strip()
            if tag and tag.lower() not in seen:
                seen.add(tag.lower())
                entries.append({"tag": tag, "category": None, "color": None})
    return _assign_palette(entries)


def _bootstrap_xlsx(data: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    rows = list(ws.iter_rows())
    if not rows:
        return []
    header = [(c.value or "").strip().lower() if isinstance(c.value, str) else "" for c in rows[0]]
    if "tag" not in header:
        return []
    ti = header.index("tag")
    seen, entries = set(), []
    for r in rows[1:]:
        if ti >= len(r):
            continue
        cell = r[ti]
        tag = (str(cell.value).strip() if cell.value is not None else "")
        if not tag or tag.lower() in seen:
            continue
        seen.add(tag.lower())
        color = None
        fill = cell.fill
        if fill and fill.fgColor and fill.fgColor.rgb and fill.patternType == "solid":
            rgb = str(fill.fgColor.rgb)
            color = rgb[-6:] if len(rgb) >= 6 else None
            if color in ("000000", "FFFFFF"):
                color = None
        entries.append({"tag": tag, "category": None, "color": color})
    return _assign_palette(entries)
