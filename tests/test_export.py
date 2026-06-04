import io

from app.core import config
from app.rag.matrix import snap_tag
from app.web.export import (
    bootstrap_codebook_from_bytes,
    to_csv,
    to_xlsx,
)


def _view():
    return {
        "view": "matrix",
        "columns": ["main_finding", "tag", "notes"],
        "rows": [
            {"source": "Doe (2020)", "year": 2020, "schema": "empiris",
             "fields": {"main_finding": "X lowers Y", "tag": "hipertensi", "notes": "(draft)"}},
            {"source": "Roe (2018)", "year": 2018, "schema": "review",
             "fields": {"main_finding": "tidak disebutkan dalam paper", "tag": "diabetes", "notes": "(draft)"}},
        ],
    }


def test_csv_export_has_header_and_rows():
    data = to_csv(_view()).decode("utf-8-sig")
    lines = [l for l in data.splitlines() if l.strip()]
    assert lines[0].startswith("source,year,schema,main_finding,tag,notes")
    assert "Doe (2020)" in lines[1]
    assert len(lines) == 3


def test_xlsx_export_colors_tag_cells():
    from openpyxl import load_workbook

    codebook = [{"tag": "hipertensi", "color": "FFD966"}, {"tag": "diabetes", "color": "9FC5E8"}]
    data = to_xlsx(_view(), codebook=codebook)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    tag_col = header.index("tag") + 1
    fill = ws.cell(row=2, column=tag_col).fill
    assert fill.patternType == "solid"
    assert fill.fgColor.rgb.endswith("FFD966")


def test_bootstrap_codebook_from_csv():
    csv_bytes = b"source,tag\nA,alpha\nB,beta\nC,alpha\n"
    entries = bootstrap_codebook_from_bytes(csv_bytes, "old.csv")
    tags = [e["tag"] for e in entries]
    assert tags == ["alpha", "beta"]  # deduped, order preserved
    assert all(e["color"] for e in entries)  # palette assigned


def test_bootstrap_codebook_from_xlsx_reads_fill_color():
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    ws.append(["source", "tag"])
    ws.append(["A", "alpha"])
    ws.cell(row=2, column=2).fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    ws.append(["B", "beta"])
    buf = io.BytesIO()
    wb.save(buf)

    entries = bootstrap_codebook_from_bytes(buf.getvalue(), "old.xlsx")
    by_tag = {e["tag"]: e for e in entries}
    assert by_tag["alpha"]["color"].endswith("FFD966")
    assert by_tag["beta"]["color"]  # palette fallback


def test_snap_tag_closed_coding():
    cb = ["Hipertensi", "Diabetes Melitus"]
    assert snap_tag("hipertensi", cb) == "Hipertensi"          # case-insensitive exact
    assert snap_tag("diabetes", cb) == "Diabetes Melitus"      # substring match
    assert snap_tag("kanker", cb).endswith("(perlu verifikasi tag)")  # unknown -> flagged
    assert snap_tag("anything", []) == "anything"              # no codebook -> unchanged


def test_export_and_codebook_endpoints(tmp_path, monkeypatch):
    import json
    from unittest.mock import patch
    import fitz
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    import app.core.db as dbmod
    import app.web.projects_routes as pr
    from app.core import projects as proj
    from app.core.db import init_db
    from app.ingest.pipeline import ingest_pdf

    def _mk(path, text):
        d = fitz.open()
        p = d.new_page()
        p.insert_text((72, 72), text, fontsize=10)
        d.save(str(path)); d.close()

    def _embed(texts, model=None, base_url=None):
        return [[0.01] * config.EMBED_DIM for _ in texts]

    db = str(tmp_path / "e.db")
    monkeypatch.setattr(dbmod, "DB_PATH", db)
    conn = init_db(db)
    a = tmp_path / "p.pdf"; _mk(a, "empirical study tag content")
    with patch("app.ingest.pipeline.embed_texts", side_effect=_embed):
        ingest_pdf(a, conn=conn)
    did = conn.execute("SELECT id FROM documents").fetchone()[0]
    pid = proj.create_project(conn, "P"); proj.add_papers(conn, pid, [did])
    conn.close()

    app = FastAPI(); app.include_router(pr.router)
    client = TestClient(app)

    # Bootstrap codebook from CSV.
    csv_bytes = b"source,tag\nA,topikA\nB,topikB\n"
    r = client.post(f"/projects/{pid}/codebook/bootstrap",
                    files={"file": ("old.csv", csv_bytes, "text/csv")})
    assert r.status_code == 200 and r.json()["learned"] == 2

    # Build matrix with closed coding (LLM tag 'topikA detail' -> snaps to 'topikA').
    def fake_chat(system, user):
        return json.dumps({"main_finding": "finding", "tag": "topikA detail", "notes": "(draft)"})

    with patch("app.rag.generator.chat", side_effect=fake_chat), \
         patch("app.rag.retriever.search", side_effect=lambda q, top_k=6, conn=None: []):
        r = client.post(f"/projects/{pid}/matrix", json={"view": "matrix"})
    assert r.status_code == 200
    assert r.json()["rows"][0]["fields"]["tag"] == "topikA"  # snapped to codebook

    # CSV export.
    r = client.get(f"/projects/{pid}/matrix/export.csv")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert b"topikA" in r.content

    # XLSX export.
    r = client.get(f"/projects/{pid}/matrix/export.xlsx")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"  # xlsx is a zip
